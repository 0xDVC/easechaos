import pytest
from datetime import datetime
from api.routes.timetable import (
    lectures_convert_to_24hour,
    exams_convert_to_24hour,
    get_json_table,
    TimeTableRequest,
)
from api.extract.extract_lectures_table import convert_to_24hour


class TestLecturesConvertTo24Hour:
    def test_morning_time(self):
        assert lectures_convert_to_24hour("7:00") == "7:00"

    def test_noon_stays_noon(self):
        assert lectures_convert_to_24hour("12:00") == "12:00"

    def test_afternoon_adds_12(self):
        assert lectures_convert_to_24hour("1:00") == "13:00"

    def test_afternoon_boundary(self):
        assert lectures_convert_to_24hour("6:59") == "18:59"

    def test_previous_pm_morning_continuation(self):
        assert lectures_convert_to_24hour("8:00", previous_was_pm=True) == "8:00"

    def test_raises_on_empty_string(self):
        with pytest.raises(ValueError, match="cannot be empty"):
            lectures_convert_to_24hour("")

    def test_raises_on_whitespace(self):
        with pytest.raises(ValueError, match="cannot be empty"):
            lectures_convert_to_24hour("   ")

    def test_raises_on_malformed_time(self):
        with pytest.raises(ValueError):
            lectures_convert_to_24hour("abc")


class TestExamsConvertTo24Hour:
    def test_am_time(self):
        assert exams_convert_to_24hour("7:00 AM") == "07:00"

    def test_pm_time(self):
        assert exams_convert_to_24hour("2:00 PM") == "14:00"

    def test_noon(self):
        assert exams_convert_to_24hour("12:00 PM") == "12:00"

    def test_midnight(self):
        assert exams_convert_to_24hour("12:00 AM") == "00:00"

    def test_morning_edge(self):
        assert exams_convert_to_24hour("11:00 AM") == "11:00"

    def test_empty_string_raises(self):
        with pytest.raises(ValueError, match="cannot be empty"):
            exams_convert_to_24hour("")


class TestGetJsonTable:
    def test_cache_hit_returns_parsed_json(self, mocker):
        mock_cache = mocker.patch("api.routes.timetable.get_table_from_cache")
        mock_cache.return_value = '[{"day": "Monday", "data": []}]'

        result = get_json_table(TimeTableRequest(class_pattern="CE 4", is_exam=False))
        assert result == [{"day": "Monday", "data": []}]

    def test_cache_miss_extracts_lecture(self, mocker, tmp_path):
        mock_cache = mocker.patch("api.routes.timetable.get_table_from_cache")
        mock_cache.return_value = None
        mock_draft = mocker.patch("api.routes.timetable._get_latest_draft")

        xlsx = tmp_path / "draft.xlsx"
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        assert(ws is not None)
        ws.title = "Monday"
        ws.append(["Classroom", "8:00-9:00", "9:00-10:00"])
        ws.append(["CE 4A", "Math 101", None])
        ws.append(["CE 4B", None, "Phy 101"])
        wb.save(xlsx)
        mock_draft.return_value = str(xlsx)

        mock_add = mocker.patch("api.routes.timetable.add_table_to_cache")

        result = get_json_table(TimeTableRequest(class_pattern="CE 4", is_exam=False))
        assert isinstance(result, list)
        assert len(result) > 0
        mock_add.assert_called_once()

    def test_cache_miss_extracts_exam(self, mocker, tmp_path):
        mock_cache = mocker.patch("api.routes.timetable.get_table_from_cache")
        mock_cache.return_value = None
        mock_draft = mocker.patch("api.routes.timetable._get_latest_draft")

        xlsx = tmp_path / "exam.xlsx"
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        assert(ws is not None)
        ws.append(["DATE", "CLASS", "PERIOD", "COURSE NAME"])
        ws.append(["45000", "CE 4A", "M", "Math 101"])
        wb.save(xlsx)
        mock_draft.return_value = str(xlsx)

        mock_add = mocker.patch("api.routes.timetable.add_table_to_cache")

        result = get_json_table(TimeTableRequest(class_pattern="CE 4", is_exam=True))
        assert isinstance(result, list)
        assert len(result) > 0
        mock_add.assert_called_once()

    def test_raises_file_not_found(self, mocker):
        mock_cache = mocker.patch("api.routes.timetable.get_table_from_cache")
        mock_cache.return_value = None
        mock_draft = mocker.patch("api.routes.timetable._get_latest_draft")
        mock_draft.return_value = "/nonexistent/path.xlsx"

        with pytest.raises(FileNotFoundError):
            get_json_table(TimeTableRequest(class_pattern="CE 4", is_exam=False))


class TestConvertTo24Hour:
    def test_hours_leq_7_become_pm(self):
        from api.extract.extract_lectures_table import convert_to_24hour as c
        assert c("7:00") == "19:00"
        assert c("1:00", is_end_time=True) == "13:00"
